from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, Form
from fastapi.responses import FileResponse
from typing import List, Optional
from pydantic import BaseModel
from datetime import datetime
import os
import shutil
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from ..models.expense import Expense
from ..schemas.expense import ExpenseCreate, ExpenseUpdate, ExpenseResponse
from ..models.receipt import Receipt
from ..models.user import User, UserRole
from ..dependencies.auth import get_current_user, require_admin
from ..services.expense_service import bulk_import_expense_rows
from ..utils.uploads import (
    validate_upload,
    EXCEL_EXTENSIONS,
    EXCEL_CONTENT_TYPES,
    MAX_EXCEL_MB,
)
import logging

logger = logging.getLogger(__name__)


class BulkImportResponse(BaseModel):
    successful_imports: int
    failed_imports: int
    errors: List[dict]

router = APIRouter()

@router.get("/", response_model=List[ExpenseResponse])
async def get_expenses(current_user: User = Depends(get_current_user)):
    # Only admins can access expenses
    if current_user.role != UserRole.ADMIN:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Only administrators can access expenses"
        )

    expenses = await Expense.find_all().to_list()

    return [
        ExpenseResponse(
            id=str(expense.id),
            expense_date=expense.expense_date,
            description=expense.description,
            amount=expense.amount,
            year=expense.year,
            month=expense.month,
            created_at=expense.id.generation_time if hasattr(expense.id, 'generation_time') else datetime.utcnow()
        )
        for expense in expenses
    ]

@router.get("/template")
async def download_expenses_template(current_user: User = Depends(require_admin)):
    template_path = os.path.join("static", "plantilla_gastos_administrativos.xlsx")
    if not os.path.exists(template_path):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Template not found")
    return FileResponse(
        path=template_path,
        filename="plantilla_gastos_administrativos.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@router.get("/{expense_id}", response_model=ExpenseResponse)
async def get_expense(expense_id: str, current_user: User = Depends(get_current_user)):
    # Only admins can access expenses
    if current_user.role != UserRole.ADMIN:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Only administrators can access expenses"
        )

    expense = await Expense.get(expense_id)
    if not expense:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Expense not found"
        )

    return ExpenseResponse(
        id=str(expense.id),
        expense_date=expense.expense_date,
        description=expense.description,
        amount=expense.amount,
        year=expense.year,
        month=expense.month,
        created_at=expense.id.generation_time if hasattr(expense.id, 'generation_time') else datetime.utcnow()
    )

@router.post("/", response_model=ExpenseResponse)
async def create_expense(
    expense_date: datetime = Form(...),
    description: str = Form(...),
    amount: float = Form(...),
    year: Optional[int] = Form(None),
    month: Optional[int] = Form(None),
    current_user: User = Depends(get_current_user)
):
    # Only admins can create expenses
    if current_user.role != UserRole.ADMIN:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Only administrators can create expenses"
        )

    # Auto-calculate year and month from expense_date if not provided
    if year is None:
        year = expense_date.year
    if month is None:
        month = expense_date.month

    expense = Expense(
        expense_date=expense_date,
        description=description,
        amount=amount,
        year=year,
        month=month
    )
    await expense.insert()

    return ExpenseResponse(
        id=str(expense.id),
        expense_date=expense.expense_date,
        description=expense.description,
        amount=expense.amount,
        year=expense.year,
        month=expense.month,
        created_at=expense.id.generation_time if hasattr(expense.id, 'generation_time') else datetime.utcnow()
    )

@router.put("/{expense_id}", response_model=ExpenseResponse)
async def update_expense(
    expense_id: str,
    expense_date: Optional[datetime] = Form(None),
    description: Optional[str] = Form(None),
    amount: Optional[float] = Form(None),
    year: Optional[int] = Form(None),
    month: Optional[int] = Form(None),
    current_user: User = Depends(get_current_user)
):
    # Only admins can update expenses
    if current_user.role != UserRole.ADMIN:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Only administrators can update expenses"
        )

    expense = await Expense.get(expense_id)
    if not expense:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Expense not found"
        )

    # Update fields
    if expense_date is not None:
        expense.expense_date = expense_date
        # If expense_date is updated and year/month are not provided, recalculate them
        if year is None:
            expense.year = expense_date.year
        if month is None:
            expense.month = expense_date.month
    if description is not None:
        expense.description = description
    if amount is not None:
        expense.amount = amount
    if year is not None:
        expense.year = year
    if month is not None:
        expense.month = month

    await expense.save()

    return ExpenseResponse(
        id=str(expense.id),
        expense_date=expense.expense_date,
        description=expense.description,
        amount=expense.amount,
        year=expense.year,
        month=expense.month,
        created_at=expense.id.generation_time if hasattr(expense.id, 'generation_time') else datetime.utcnow()
    )

@router.delete("/{expense_id}")
async def delete_expense(expense_id: str, current_user: User = Depends(get_current_user)):
    # Only admins can delete expenses
    if current_user.role != UserRole.ADMIN:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Only administrators can delete expenses"
        )

    expense = await Expense.get(expense_id)
    if not expense:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Expense not found"
        )
    await expense.delete()
    return {"message": "Expense deleted successfully"}

# Expenses don't generate receipts - they receive vouchers from suppliers
# This endpoint is no longer needed


@router.post("/bulk-import", response_model=BulkImportResponse)
async def bulk_import_expenses(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user)
):
    """Bulk import expenses from Excel file"""
    if current_user.role != UserRole.ADMIN:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Only administrators can bulk import expenses"
        )

    await validate_upload(file, EXCEL_EXTENSIONS, EXCEL_CONTENT_TYPES, MAX_EXCEL_MB)

    try:
        # Load workbook
        wb = load_workbook(file.file)
        ws = wb.active

        # Expected columns: Fecha de Gasto, Descripción, Monto, Año, Mes
        expected_headers = ['Fecha de Gasto', 'Descripción', 'Monto', 'Año', 'Mes']

        # Check headers
        actual_headers = []
        for col in range(1, 6):
            cell_value = ws.cell(row=1, column=col).value
            if cell_value:
                actual_headers.append(str(cell_value).strip())
            else:
                actual_headers.append("")

        if actual_headers != expected_headers:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Invalid Excel format. Expected headers: {', '.join(expected_headers)}"
            )

        results = await bulk_import_expense_rows(ws)
        return BulkImportResponse(**results)

    except InvalidFileException:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Invalid Excel file"
        )
    except Exception:
        logger.exception("Error processing expenses import file")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Error processing file"
        )